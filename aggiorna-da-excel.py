#!/usr/bin/env python3
"""
aggiorna-da-excel.py
---------------------
Alternativa "senza codice" ad aggiorna-recensioni.py: legge le recensioni
direttamente dal foglio Excel dati/risultati-form.xlsx (foglio "Recensioni"),
tiene solo quelle con "Pubblica" = SI, rigenera dati/recensioni-generali.json
(e il suo specchio dati/recensioni-generali.js, letto dalla homepage) e poi
rigenera la griglia recensioni dentro recensioni.html.

COME FUNZIONA IN PRATICA
  1. Apri dati/risultati-form.xlsx, foglio "Recensioni".
  2. Aggiungi/modifica righe: nome, ruolo, voto, testo, anno...
  3. Nella colonna "Pubblica (SI/NO)" scrivi SI per le righe che vuoi
     mostrare sul sito, nella pagina Recensioni generali (NO o vuoto per
     tenerle solo in archivio).
  4. Nella colonna "In home (SI/NO)" scrivi SI per le righe che, oltre a
     comparire in Recensioni generali, vuoi mostrare anche in homepage
     (la stessa cosa che puoi fare anche dal pannello admin, sezione
     Recensioni — questa colonna e il pannello scrivono lo stesso dato).
  5. Salva il file Excel.
  6. Lancia:  python3 aggiorna-da-excel.py
  7. Carica i file aggiornati (recensioni.html, dati/recensioni-generali.json
     e dati/recensioni-generali.js) sul sito — o pubblica dal pannello admin.

USO:
    python3 aggiorna-da-excel.py --excel dati/risultati-form.xlsx --progetto .
"""

import argparse
import json
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit(
        "Manca il pacchetto 'openpyxl'. Installalo con:\n"
        "    pip install openpyxl --break-system-packages\n"
        "(oppure: pip install openpyxl)"
    )

MARCATORE_INIZIO = "<!-- INIZIO GRIGLIA GENERATA -->"
MARCATORE_FINE = "<!-- FINE GRIGLIA GENERATA -->"

# Intestazioni attese nel foglio "Recensioni" (riga 1)
COLONNE = ["pubblica", "in_home", "nome", "ruolo", "email", "voto", "testo", "anno", "data", "note"]


def leggi_recensioni_da_excel(percorso_excel: Path, esistenti: dict) -> list:
    wb = openpyxl.load_workbook(percorso_excel, data_only=True)
    if "Recensioni" not in wb.sheetnames:
        raise SystemExit(f"Non trovo il foglio 'Recensioni' in {percorso_excel}")
    ws = wb["Recensioni"]

    righe = list(ws.iter_rows(min_row=2, values_only=True))
    recensioni = []
    for riga in righe:
        if riga is None or all(v is None or str(v).strip() == "" for v in riga):
            continue  # riga vuota

        valori = dict(zip(COLONNE, riga))

        pubblica = str(valori.get("pubblica") or "").strip().upper()
        if pubblica != "SI" and pubblica != "SÌ":
            continue  # non pubblicata: saltata

        nome = str(valori.get("nome") or "").strip()
        testo = str(valori.get("testo") or "").strip()
        if not nome or not testo or "(ESEMPIO)" in nome.upper():
            continue  # riga incompleta o riga di esempio: saltata

        in_home = str(valori.get("in_home") or "").strip().upper()
        pinned = in_home in ("SI", "SÌ")

        try:
            voto = int(float(valori.get("voto") or 5))
        except (TypeError, ValueError):
            voto = 5
        voto = max(1, min(5, voto))

        prec = esistenti.get(nome.strip().lower(), {})
        recensioni.append({
            "nome": nome,
            "ruolo": str(valori.get("ruolo") or "").strip(),
            "testo": testo,
            "voto": voto,
            "anno": str(valori.get("anno") or "").strip() or "2026",
            "pinned": pinned,
            # Foto e "categoria" curate a mano (es. dal pannello admin, o in
            # precedenza in dati/recensioni-generali.json) non sono in questo
            # Excel: le conserviamo se la stessa persona le aveva già.
            "avatar": prec.get("avatar"),
            "categoria": prec.get("categoria"),
        })

    return recensioni


def blocco_card(rec: dict) -> str:
    voto = int(rec.get("voto", 5))
    voto = max(1, min(5, voto))
    stelle = "★" * voto + "☆" * (5 - voto)
    ruolo = rec.get("ruolo", "").strip()
    return f"""          <article class="review-card reveal">

            <div class="review-header">

              <div class="review-avatar review-avatar-lettera" aria-hidden="true">
                {rec['nome'].strip()[:1].upper()}
              </div>

              <div class="review-user">
                <h3>{rec['nome'].strip()}</h3>
                <p>{ruolo}</p>
              </div>

            </div>

            <div class="review-rating" aria-label="{voto} stelle su 5">
              {stelle}
            </div>

            <blockquote>
              &ldquo;{rec['testo'].strip()}&rdquo;
            </blockquote>

            <div class="review-footer">
              <span>{rec.get('categoria') or 'Recensione generale'}</span>
              <span>{rec.get('anno', '')}</span>
            </div>

          </article>
"""


def blocco_vuoto() -> str:
    return """          <div class="reviews-empty reveal">
            <p>Ancora nessuna recensione generale pubblicata.<br>Sii il primo a lasciarne una, dal modulo qui sotto.</p>
          </div>
"""


def genera_griglia(recensioni: list) -> str:
    if not recensioni:
        return blocco_vuoto()
    return "\n".join(blocco_card(r) for r in recensioni)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--progetto", default=".")
    p.add_argument("--excel", default="dati/risultati-form.xlsx")
    p.add_argument("--dati", default="dati/recensioni-generali.json")
    a = p.parse_args()

    progetto = Path(a.progetto)
    excel_path = progetto / a.excel
    if not excel_path.exists():
        raise SystemExit(f"Non trovo il file Excel: {excel_path}")

    dati_path = progetto / a.dati
    esistenti = {}
    if dati_path.exists():
        try:
            for r in json.loads(dati_path.read_text(encoding="utf-8")):
                nome = str(r.get("nome") or "").strip().lower()
                if nome:
                    esistenti[nome] = r
        except (json.JSONDecodeError, OSError):
            pass

    recensioni = leggi_recensioni_da_excel(excel_path, esistenti)

    dati_path.parent.mkdir(parents=True, exist_ok=True)
    dati_path.write_text(
        json.dumps(recensioni, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"{dati_path} aggiornato — {len(recensioni)} recensioni pubblicate")

    js_path = dati_path.with_suffix(".js")
    js_path.write_text(
        "window.GPJ_REVIEWS = " + json.dumps(recensioni, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"{js_path} aggiornato (usato dalla homepage per le recensioni pinnate)")

    pagina = progetto / "recensioni.html"
    testo = pagina.read_text(encoding="utf-8")

    if MARCATORE_INIZIO not in testo or MARCATORE_FINE not in testo:
        raise SystemExit(
            "Non trovo i marcatori della griglia in recensioni.html. "
            "Non tocco il file per sicurezza."
        )

    griglia = genera_griglia(recensioni)
    nuovo = re.sub(
        re.escape(MARCATORE_INIZIO) + r".*?" + re.escape(MARCATORE_FINE),
        MARCATORE_INIZIO + "\n\n" + griglia + "\n        " + MARCATORE_FINE,
        testo, count=1, flags=re.DOTALL)

    pagina.write_text(nuovo, encoding="utf-8")
    print(f"recensioni.html rigenerata — {len(recensioni)} recensioni generali pubblicate")


if __name__ == "__main__":
    main()
